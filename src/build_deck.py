#!/usr/bin/env python3
"""Agent Skills training deck — MSCI Creative Mode (single prs, single save)."""
import sys, os
OUT = os.environ.get("BUNDLE_OUT", "/mnt/user-data/outputs")
SKILL = "/mnt/skills/plugins/msci-pptx-skill"
sys.path.insert(0, os.path.join(SKILL, "scripts"))

from pptx import Presentation
import creative_builders as cb
from creative_builders import (
    add_canvas_slide, add_cover_slide, add_divider_slide, add_closing_slide,
    add_disclaimer_slide, detect_layout_indices, delete_all_slides, fix_shape_ids,
    auto_style,
    build_icon_grid, build_walled_garden, build_quadrant_matrix, build_styled_table,
    build_horizontal_bar_stack, build_split_panel_stats, build_accordion_rows,
    build_checklist_rows, build_data_flow_diagram, build_honeycomb_grid,
    build_three_pillar, build_terminal_code_block, build_side_by_side_projects,
    build_kanban_columns, build_traffic_light_grid, build_three_containers_bar,
    build_numbered_step_flow, build_big_number_hero, build_exec_summary_text,
    build_connected_app_cards, build_shadow_cards_badges, build_horizontal_card_row,
    build_progress_tracker, build_nested_layers, build_exploded_container,
    build_quote_hero, build_inverted_funnel,
)
from creative_builders import BLUE, TEAL, SMOKEY, SKY, GREEN, ICE, PALE_BLUE, PALE_TEAL

SRC = "Source: DeepLearning.AI 'Agent Skills with Anthropic'; agentskills.io specification; Anthropic platform documentation, Sept 2026"
S = auto_style("training")

prs = Presentation(os.path.join(SKILL, "MSCI_Creative_Template.pptx"))
detect_layout_indices(prs)
delete_all_slides(prs)

DIVS = [cb.LAYOUT_DIVIDER_TEAL, cb.LAYOUT_DIVIDER_BLUE, cb.LAYOUT_DIVIDER_IMAGE]
_d = {"i": 0}
def divider(title, subtitle, label):
    lay = DIVS[_d["i"] % len(DIVS)]
    _d["i"] += 1
    return add_divider_slide(prs, title=title, subtitle=subtitle, section_label=label, layout_index=lay)


# ============================================================ cover
add_cover_slide(
    prs,
    title="Agent Skills",
    subtitle="Packaging repeatable workflows so an agent works the way your team does",
    presenter_date="Training session  ·  September 2026",
)

# ============================================================ 1 · Foundations
divider("Foundations", "What a skill is, and when a prompt has become one", "Section 01")

sl = add_canvas_slide(prs, "A skill is a folder of instructions — that is the whole idea",
                      section="Foundations", source=SRC)
build_icon_grid(sl, [
    {"icon_name": "24_document", "title": "SKILL.md",
     "description": "The only required file. Frontmatter plus a markdown body."},
    {"icon_name": "37_code", "title": "scripts/",
     "description": "Code the agent executes. Only the output enters context."},
    {"icon_name": "29_bookshelf", "title": "references/",
     "description": "Documents read on demand, when the body asks for them."},
    {"icon_name": "05_box_closed", "title": "assets/",
     "description": "Templates, logos and schemas used in the output."},
], accent_color=BLUE)

sl = add_canvas_slide(prs, "The trigger is repetition, not complexity",
                      section="Foundations", source=SRC)
build_walled_garden(sl,
    left_panel={"title": "Every week, by hand",
                "subtitle": "The prompt-based workflow",
                "items": ["Retype or re-paste the specification",
                          "Rely on knowledge the operator may not have",
                          "Every pasted document stays in context",
                          "Lives in one person's chat history"],
                "color": SMOKEY, "bg_color": ICE},
    right_panel={"title": "Once, as a folder",
                 "subtitle": "The same workflow as a skill",
                 "items": ["Instructions live in SKILL.md",
                           "Conditional rules in references/",
                           "Loaded only when the request matches",
                           "Shareable, editable, versioned"],
                 "color": TEAL, "bg_color": PALE_TEAL},
    accent_color=BLUE, style=S["wg"])

sl = add_canvas_slide(prs, "Four costs a skill removes, and where each one goes",
                      section="Foundations", source=SRC)
build_horizontal_card_row(sl, [
    {"title": "Repetition", "badge": "COST 1",
     "description": "The specification moves into the SKILL.md body and stops being retyped every week."},
    {"title": "Tacit knowledge", "badge": "COST 2",
     "description": "Benchmarks and decision rules get written down where anyone can read and review them."},
    {"title": "Context pollution", "badge": "COST 3",
     "description": "Bulky conditional material moves to references/ and loads only when it is asked for."},
    {"title": "No distribution", "badge": "COST 4",
     "description": "A folder can be zipped, shared, installed elsewhere and kept under version control."},
], accent_color=BLUE, style="tinted")

sl = add_canvas_slide(prs, "Frontmatter is the only part read before anything loads",
                      section="Foundations", source="Source: agentskills.io specification, published December 2025")
build_styled_table(sl,
    ["Field", "Required", "Limit", "What it is for"],
    [["name", "Yes", "64 chars", "Lowercase, hyphens; must match the folder name"],
     ["description", "Yes", "1,024 chars", "What it does and when to use it — the routing key"],
     ["license", "No", "—", "License name or a bundled license file"],
     ["compatibility", "No", "500 chars", "Environment needs: packages, network, host"],
     ["metadata", "No", "—", "Arbitrary key-value pairs such as author or version"],
     ["allowed-tools", "No", "Experimental", "Pre-approved tools; honoured in some hosts only"]],
    accent_color=BLUE)

# ============================================================ 2 · Progressive disclosure
divider("Progressive disclosure", "Why a hundred skills cost almost nothing until they fire", "Section 02")

sl = add_canvas_slide(prs, "Three tiers, and only the first one is always paid for",
                      section="Progressive disclosure", source=SRC)
build_horizontal_bar_stack(sl, [
    {"title": "Tier 1 — Metadata", "width_pct": 1.0,
     "description": "Name and description for every installed skill. Loaded at session start."},
    {"title": "Tier 2 — Instructions", "width_pct": 0.72,
     "description": "The full SKILL.md body, loaded when the description matches the request."},
    {"title": "Tier 3 — Resources", "width_pct": 0.44,
     "description": "Reference files, assets and scripts, only when the body directs it."},
], accent_color=BLUE)

sl = add_canvas_slide(prs, "The cost profile is what should drive your design",
                      section="Progressive disclosure", source=SRC)
build_split_panel_stats(sl,
    headline="Context is a shared resource",
    body_text="Every token you add raises cost, fills the window faster and increases the chance of a degraded answer. Progressive disclosure is an accuracy mechanism as much as an economy one.",
    stats=[{"number": "~100", "label": "tokens per installed skill, always loaded"},
           {"number": "5,000", "label": "token target for the SKILL.md body"},
           {"number": "70-90%", "label": "reported saving versus static context"},
           {"number": "100+", "label": "skills a single system can carry"}],
    accent_color=BLUE)

sl = add_canvas_slide(prs, "Three places progressive disclosure does not save you",
                      section="Progressive disclosure", source=SRC)
build_accordion_rows(sl, [
    {"header": "A skill attached to a subagent",
     "detail": "The full SKILL.md is injected when the subagent is dispatched. No further disclosure happens, so budget for the whole body."},
    {"header": "A reference file read unconditionally",
     "detail": "If the body always reads it, you have written a long SKILL.md in two files. The saving comes from the condition, not the split."},
    {"header": "A bloated skill library",
     "detail": "Metadata cost is per installed skill, not per used skill. Fifty unused skills still carry a standing tax. Prune them."},
], accent_color=BLUE)

# ============================================================ 3 · Placement
divider("Placement", "Skills, tools, MCP, subagents and prompts", "Section 03")

sl = add_canvas_slide(prs, "Five primitives, distinguished by what they cost when idle",
                      section="Placement", source=SRC)
build_styled_table(sl,
    ["Primitive", "What it gives you", "Cost when not in use"],
    [["Prompt", "One-time instruction in this conversation", "Nothing — and nothing persists either"],
     ["Tool", "A raw capability the agent can call", "Its definition sits in context permanently"],
     ["MCP server", "Connectivity to external systems and data", "Its tool definitions sit in context permanently"],
     ["Skill", "Procedural knowledge: which tools, what order", "About 100 tokens of metadata"],
     ["Subagent", "An isolated window and restricted tools", "Nothing until dispatched, then a whole window"]],
    accent_color=BLUE)

sl = add_canvas_slide(prs, "Choosing the primitive is a context and permission decision",
                      section="Placement", source=SRC)
build_checklist_rows(sl, [
    {"title": "Reach an external system or dataset", "detail": "MCP server — it brings the tools, not the method", "status": "blue"},
    {"title": "Encode our house method for a recurring job", "detail": "Skill — and keep the body a sequence, not a rulebook", "status": "green"},
    {"title": "Ship the same expertise to many agents", "detail": "Skill, not a subagent. Subagents isolate; skills distribute", "status": "green"},
    {"title": "Isolate a noisy task from the main window", "detail": "Subagent — its transcript never reaches the parent", "status": "amber"},
    {"title": "Restrict what a step may touch", "detail": "Subagent tool allow-list. Prose in a SKILL.md is not a control", "status": "red"},
    {"title": "One-off instruction for this turn", "detail": "A prompt. Not everything needs to be an asset", "status": "gray"},
], accent_color=BLUE)

sl = add_canvas_slide(prs, "How they compose in a real workflow",
                      section="Placement", source=SRC)
build_data_flow_diagram(sl,
    source={"title": "MCP brings the data", "description": "Warehouses, drives, ticketing and document stores, reached through tools the agent already has.",
            "icon_name": "07_map_pin"},
    process={"title": "Subagents parallelise", "description": "Independent investigations run in isolated windows and report only their findings back.",
             "icon_name": "34_cycle"},
    output={"title": "Skills make it predictable", "description": "The same structure, the same order and the same output contract on every run.",
            "icon_name": "15_checklist"},
    accent_color=BLUE)

# ============================================================ 4 · What ships
divider("What ships in the box", "Document skills, example skills and skill-creator", "Section 04")

sl = add_canvas_slide(prs, "Two families of skills ship from the vendor repository",
                      section="What ships in the box", source="Source: github.com/anthropics/skills")
build_honeycomb_grid(sl, [
    {"title": "docx", "category": "Document skill"},
    {"title": "pdf", "category": "Document skill"},
    {"title": "pptx", "category": "Document skill"},
    {"title": "xlsx", "category": "Document skill"},
    {"title": "skill-creator", "category": "Example — on by default"},
    {"title": "mcp-builder", "category": "Example skill"},
    {"title": "webapp-testing", "category": "Example skill"},
    {"title": "brand-guidelines", "category": "Example skill"},
    {"title": "frontend-design", "category": "Example skill"},
    {"title": "canvas-design", "category": "Example skill"},
    {"title": "doc-coauthoring", "category": "Example skill"},
    {"title": "algorithmic-art", "category": "Example skill"},
], accent_color=BLUE)

sl = add_canvas_slide(prs, "Three separate stores — installing once is not installing everywhere",
                      section="What ships in the box", source=SRC)
build_three_pillar(sl,
    base_title="A skill uploaded in Claude.ai is invisible to the API and to Claude Code",
    base_desc="Same folder, three installs. This is the single most common source of confusion when moving from the web app to code.",
    pillars=[
        {"title": "Claude.ai and Desktop",
         "description": "Document skills built in and always on. Example skills toggleable, off by default except skill-creator. Custom skills uploaded as a zip."},
        {"title": "Claude Code",
         "description": "Ships with none. Add the marketplace, install the collections, then restart. Project skills in .claude/skills, personal in ~/.claude/skills."},
        {"title": "API and Agent SDK",
         "description": "API: upload the skill, then mount it in container.skills, maximum eight. SDK: on disk, plus the Skill tool and setting_sources."},
    ], accent_color=BLUE)

sl = add_canvas_slide(prs, "skill-creator does far more than scaffold a folder",
                      section="What ships in the box", source="Source: anthropics/skills, skill-creator")
build_terminal_code_block(sl,
    code_text=("# package a finished skill\n"
               "python -m scripts.package_skill ./my-skill\n\n"
               "# benchmark one iteration against a no-skill baseline\n"
               "python -m scripts.aggregate_benchmark workspace/iteration-1 \\\n"
               "    --skill-name my-skill\n\n"
               "# optimise the description on a held-out split\n"
               "python -m scripts.run_loop --eval-set evals/evals.json \\\n"
               "    --skill-path ./my-skill --max-iterations 5"),
    tip_cards=[
        {"title": "Grades your work", "description": "It carries the best-practice list in a form that can score a skill you already wrote."},
        {"title": "Measures the delta", "description": "Pass rate, tokens and duration, with and without the skill."},
        {"title": "Tunes the description", "description": "Three repeats per query, 60/40 split, selection on held-out items."},
    ], accent_color=BLUE)

# ============================================================ 5 · Authoring
divider("Authoring", "Descriptions, degrees of freedom and where payloads go", "Section 05")

sl = add_canvas_slide(prs, "The description is the only thing read before the skill loads",
                      section="Authoring", source=SRC)
build_side_by_side_projects(sl,
    left_project={"title": "Will not fire reliably",
                  "instruction": "Vague scope, no trigger words, first person",
                  "prompt": "\"Helps with analysing data and producing useful summaries for the team.\"",
                  "icon_name": "27_faq"},
    right_project={"title": "Fires when it should",
                   "instruction": "States what and when, in the third person, with the words a user would type",
                   "prompt": "\"Analyze weekly marketing campaign performance from CSV or BigQuery data. Use when the user asks about campaign metrics, funnel analysis, ROAS, or budget reallocation.\"",
                   "icon_name": "15_checklist"},
    accent_color=BLUE)

sl = add_canvas_slide(prs, "Set the freedom dial by asking whether variation is a defect",
                      section="Authoring", source=SRC)
build_kanban_columns(sl, [
    {"title": "Low freedom", "cards": [
        "Compliance and audit output",
        "Numbered, mandatory sequence",
        "Fixed template in assets/",
        "Explicit reason a step may be skipped",
        "Named output tree"]},
    {"title": "Medium freedom", "cards": [
        "Analysis and reporting",
        "Fixed order, flexible wording",
        "Deterministic checks in scripts/",
        "Interpretation left to the model"]},
    {"title": "High freedom", "cards": [
        "Concept and design work",
        "Goals and constraints, not steps",
        "Palette ranges, not hex values",
        "Guardrails instead of scripts",
        "Many acceptable outputs"]},
], accent_color=BLUE)

sl = add_canvas_slide(prs, "Filing a payload is really a token-budgeting decision",
                      section="Authoring", source=SRC)
build_traffic_light_grid(sl,
    row_labels=["scripts/", "references/", "assets/"],
    col_labels=["Loading behaviour", "Context cost", "Best for"],
    cells=[
        [{"text": "Executed", "status": "green"},
         {"text": "Output only", "status": "green"},
         {"text": "Deterministic checks and file mechanics", "status": "blue"}],
        [{"text": "Read on demand", "status": "amber"},
         {"text": "Full length at read time", "status": "amber"},
         {"text": "Conditional and bulky material", "status": "blue"}],
        [{"text": "Consumed by code", "status": "green"},
         {"text": "Often nothing", "status": "green"},
         {"text": "Templates, schemas, brand files", "status": "blue"}],
    ], accent_color=BLUE)

sl = add_canvas_slide(prs, "The most transplantable architecture in the course",
                      section="Authoring", source=SRC)
build_three_containers_bar(sl,
    containers=[
        {"title": "scripts/ for determinism",
         "description": "The statistical work lives in Python, so identical input yields identical numbers and only stdout costs tokens."},
        {"title": "references/ for interpretation",
         "description": "How to read a stationarity or autocorrelation result, opened only when the agent must explain a number."},
        {"title": "assets/ for contracts",
         "description": "Output templates and schemas, one per format, so the unused one never loads."},
    ],
    shared_bar_title="SKILL.md sequences them — and does nothing else",
    shared_bar_desc="Which script, in what order, what may be skipped and why, and the exact output tree the run must produce.",
    accent_color=BLUE)

# ============================================================ 6 · Harnesses
divider("Four harnesses", "The same folder in Claude.ai, the API, Claude Code and the SDK", "Section 06")

sl = add_canvas_slide(prs, "On the API, a skills request has five moving parts",
                      section="Four harnesses", source="Source: Anthropic platform documentation, reviewed September 2026")
build_numbered_step_flow(sl, [
    {"title": "Model", "description": "Any current model. Skills are not a model feature."},
    {"title": "Betas", "description": "Skills, code execution, files. The skills header is optional post-GA."},
    {"title": "Skills list", "description": "container.skills names type, id and version. Eight maximum."},
    {"title": "Code exec tool", "description": "Mandatory. Supplies the container, shell and filesystem."},
    {"title": "Files API", "description": "Upload inputs by file_id, download the artefacts."},
], accent_color=BLUE, style=S["nsf"])

sl = add_canvas_slide(prs, "The API sandbox is stricter than the one Claude.ai gives you",
                      section="Four harnesses", source="Source: Anthropic code execution tool documentation")
build_big_number_hero(sl, [
    {"number": "8", "label": "skills per container"},
    {"number": "5 GiB", "label": "RAM and workspace disk, 1 CPU"},
    {"number": "0", "label": "network access — no runtime pip install"},
    {"number": "500 MB", "label": "maximum per uploaded file"},
], accent_color=BLUE)

sl = add_canvas_slide(prs, "In Claude Code, decide what is always in context and what is not",
                      section="Four harnesses", source=SRC)
build_exec_summary_text(sl, [
    {"heading": "CLAUDE.md — always loaded",
     "body": "Stack, architecture, file layout and the conventions that apply on every turn. If a rule is useful on a turn unrelated to the task, it belongs here."},
    {"heading": ".claude/skills/ — on demand",
     "body": "One folder per skill. Rules that apply only while doing one class of work: how to add a command, how to write its tests, how to review it. Restart to pick up new skills."},
    {"heading": ".claude/agents/ — on dispatch",
     "body": "One markdown file per subagent, with frontmatter naming its tools, model and skills. Nothing is inherited from the parent, so both are explicit."},
], accent_color=BLUE)

sl = add_canvas_slide(prs, "Split the development loop so the noise never reaches the parent",
                      section="Four harnesses", source=SRC)
build_connected_app_cards(sl, [
    {"title": "code-reviewer", "verb": "REVIEW",
     "prompt": "Bash Glob Grep Read — read-only on purpose, so it reports issues instead of quietly fixing them.",
     "icon_name": "17_magnifier"},
    {"title": "test-generator-runner", "verb": "TEST",
     "prompt": "Adds Edit and Write, because it must create test files, then runs the suite and returns a summary.",
     "icon_name": "15_checklist"},
    {"title": "parent agent", "verb": "BUILD",
     "prompt": "Stays on development and receives only verdicts and diffs. Its window never fills with run transcripts.",
     "icon_name": "26_developer"},
], ribbon_title="Each subagent carries its own skill — nothing is inherited", accent_color=BLUE)

sl = add_canvas_slide(prs, "In the Agent SDK, four settings decide whether skills exist at all",
                      section="Four harnesses", source="Source: Claude Agent SDK documentation, reviewed September 2026")
build_shadow_cards_badges(sl, [
    {"title": "setting_sources", "badge": "REQUIRED",
     "description": "Must include 'project' or 'user'. Omit it and no skill loads, with no error message."},
    {"title": "Skill in allowed_tools", "badge": "REQUIRED",
     "description": "Without it skills are invisible, even when the folder is correct and present."},
    {"title": "Task in allowed_tools", "badge": "FOR SUBAGENTS",
     "description": "The dispatch tool. Declared agents can never be invoked without it."},
    {"title": "Explicit tool grants", "badge": "NOT DEFAULT",
     "description": "Read, Grep and Glob are allowed by default. Write, Bash, WebSearch and WebFetch are not."},
], accent_color=BLUE, style=S["scb"])

# ============================================================ 7 · Evaluate, secure, ship
divider("Evaluate, secure, ship", "Turning a skill into something you can defend", "Section 07")

sl = add_canvas_slide(prs, "Three failure modes, and they need measuring separately",
                      section="Evaluate, secure, ship", source=SRC)
build_horizontal_card_row(sl, [
    {"title": "It did not fire", "badge": "ROUTING",
     "description": "A description problem. Measure trigger rate over repeated runs, and include queries that should not fire it."},
    {"title": "It fired and did the wrong thing", "badge": "BEHAVIOUR",
     "description": "Assert things a script could check: file at path, sections present, steps executed in the stated order."},
    {"title": "It fired and cost too much", "badge": "ECONOMY",
     "description": "Compare tokens and duration against a baseline run with the skill disabled. Value is the delta."},
], accent_color=BLUE, style=S["hcr"])

sl = add_canvas_slide(prs, "A credible evaluation loop, in five steps",
                      section="Evaluate, secure, ship", source=SRC)
build_progress_tracker(sl, [
    {"title": "Write the eval set", "time": "once",
     "description": "Should-fire and should-not-fire queries, plus objectively checkable behavioural assertions."},
    {"title": "Run with the skill", "time": "x3",
     "description": "Repeat every query at least three times, because triggering is stochastic."},
    {"title": "Run the baseline", "time": "x3",
     "description": "The same queries with the skill disabled. Without this you are measuring the model, not the skill."},
    {"title": "Aggregate", "time": "per run",
     "description": "Pass rate, tokens and duration as means with standard deviations, and the delta against the last iteration."},
    {"title": "Close the gates", "time": "before ship",
     "description": "Human review, then a pass on every model you deploy on. Regression-test consistency of structure."},
], accent_color=BLUE)

sl = add_canvas_slide(prs, "SKILL.md is the least dangerous file in the folder",
                      section="Evaluate, secure, ship", source=SRC)
build_inverted_funnel(sl, [
    {"title": "Description — everyone reads it", "description": "The only part shown in a settings list, and all most reviewers ever see."},
    {"title": "SKILL.md body — usually skimmed", "description": "May instruct the agent to read or execute anything further down."},
    {"title": "references/ and assets/ — read by the agent", "description": "Often by nobody else. Templates and images are content too."},
    {"title": "scripts/ — read by almost nobody", "description": "Executable code: what it runs, writes, and sends off the machine."},
], accent_color=BLUE)

sl = add_canvas_slide(prs, "Ten minutes of audit, before you enable anything",
                      section="Evaluate, secure, ship", source=SRC)
build_checklist_rows(sl, [
    {"title": "Source is trusted", "detail": "An internal registry, or a vendor repository you verified yourself", "status": "green"},
    {"title": "Every script read", "detail": "What it executes, what it writes, what it sends outside the machine", "status": "red"},
    {"title": "No external fetch-then-act", "detail": "The top red flag: it delegates the skill's contents to whoever controls the URL", "status": "red"},
    {"title": "No secrets embedded", "detail": "Skills get zipped, shared and installed elsewhere. Anything inside travels", "status": "amber"},
    {"title": "Tool surface acceptable", "detail": "Does it imply Bash, Write or WebFetch? Enforcement is the allow-list, not the prose", "status": "amber"},
    {"title": "Version pinned, owner named", "detail": "An unowned skill encoding a methodology drifts into confidently wrong output", "status": "blue"},
], accent_color=BLUE)

sl = add_canvas_slide(prs, "Write to the safe column; date-stamp the rest",
                      section="Evaluate, secure, ship", source="Source: agentskills.io client showcase and specification")
build_exec_summary_text(sl, [
    {"heading": "Safe in roughly 40 hosts",
     "body": "Name and description, a plain markdown body, and scripts, references and assets reached by relative path with forward slashes. Claude Code, Codex, Copilot, VS Code, Cursor and Gemini CLI all read this."},
    {"heading": "Host-dependent — verify in your version",
     "body": "The experimental allowed-tools field, the subagent skills field, marketplace installation, and the API's container.skills parameter. Each works somewhere and is ignored or absent elsewhere."},
    {"heading": "Never assume",
     "body": "How the skill is surfaced to the user, whether the host has network access, which packages are pre-installed, or that a script may install its own dependencies at runtime."},
], accent_color=BLUE)

# ============================================================ 8 · Applied
divider("Applied", "What this looks like on a real verification workflow", "Section 08")

sl = add_canvas_slide(prs, "A batch verification pipeline is an almost perfect fit",
                      section="Applied", source=SRC)
build_exploded_container(sl,
    center_title="SKILL.md\nsequence, skip rules,\noutput contract",
    components=[
        {"title": "scripts/", "description": "Schema validation, the mechanical checks, and summarisation into a fixed results tree."},
        {"title": "references/", "description": "Methodology and per-flag decision rules, opened only when a case needs adjudicating."},
        {"title": "assets/", "description": "Column contract and standard remark phrasing, consumed by code rather than read."},
        {"title": "results/", "description": "checks.json, summary.txt, residuals.csv — named paths are what make a run auditable later."},
    ], accent_color=BLUE)

sl = add_canvas_slide(prs, "And the case for not building one",
                      section="Applied", source=SRC)
build_quote_hero(sl,
    quote_text="The threshold is repetition plus a settled convention — not complexity. A workflow that runs twice, or whose rules are still being argued about, is better served by a prompt in a shared document than by a folder someone now has to own.",
    attribution="When a skill is premature",
    accent_color=BLUE)

# ============================================================ close
close = add_closing_slide(prs, title="Build one this week",
                  subtitle="Take the workflow you have retyped most often, write the description first, and run it against a baseline before you trust it.")
for _ph in list(close.placeholders):
    if _ph.placeholder_format.idx == 10:
        cb._blank_ph(_ph)
add_disclaimer_slide(prs)

fix_shape_ids(prs)
out = os.path.join(OUT, "Agent-Skills-Training-Deck.pptx")
prs.save(out)
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst), "saved:", out)
